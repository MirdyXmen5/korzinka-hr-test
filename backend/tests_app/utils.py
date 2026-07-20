import openpyxl
from .models import Test, Question

CYRILLIC_TO_LATIN = {'А': 'A', 'а': 'A', 'Б': 'B', 'б': 'B', 'В': 'C', 'в': 'C'}

def normalize_answers(raw_answer):
    if not raw_answer:
        return 'A'
    parts = [p.strip() for p in str(raw_answer).split(',')]
    normalized = []
    for p in parts:
        mapped = CYRILLIC_TO_LATIN.get(p, p.upper())
        if mapped in ('A', 'B', 'C'):
            normalized.append(mapped)
    return ','.join(normalized) if normalized else 'A'

def parse_excel(file, test_name=None, language=None):
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    created_tests = []
    
    if test_name and language:
        ws = wb.active
        test = _parse_sheet(ws, test_name, language)
        created_tests.append(test)
    else:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lang = 'kk' if '(каз)' in sheet_name.lower() or '(қаз)' in sheet_name.lower() else 'ru'
            name = test_name or sheet_name.replace('(каз)', '').replace('(қаз)', '').strip()
            test = _parse_sheet(ws, name if test_name else sheet_name, lang)
            created_tests.append(test)
    
    wb.close()
    return created_tests

def _parse_sheet(ws, name, language):
    test = Test.objects.create(name=name, language=language)
    questions = []
    
    for idx, row in enumerate(ws.iter_rows(min_row=2, max_col=5, values_only=True)):
        question_text = row[0]
        if not question_text:
            continue
        
        option_a = row[1] or ''
        option_b = row[2] or ''
        option_c = row[3] or ''
        raw_correct = row[4] if len(row) > 4 else None
        correct_answers = normalize_answers(raw_correct)
        
        questions.append(Question(
            test=test,
            text=str(question_text).strip(),
            option_a=str(option_a).strip(),
            option_b=str(option_b).strip(),
            option_c=str(option_c).strip(),
            correct_answers=correct_answers,
            order=idx + 1
        ))
    
    Question.objects.bulk_create(questions)
    return test

def parse_excel_from_path(file_path, test_name=None, language=None):
    return parse_excel(file_path, test_name, language)
